import openpyxl
import os
import sys
import requests
import time

sys.stdout.reconfigure(encoding='utf-8')

SUPABASE_URL = "https://cktrjbcuvopkxsjigjru.supabase.co"
SUPABASE_KEY = "sb_publishable_814KmbPEV-6zprnxIyoxnA_FbYtdVAJ"
EXCEL_FILE = "Costaff_Master Tracker_V.1.xlsx"

HEADERS = {
    "apikey": SUPABASE_KEY,
    "Authorization": f"Bearer {SUPABASE_KEY}",
    "Content-Type": "application/json",
    "Prefer": "return=minimal"
}

def post_batch(table, records):
    if not records:
        return 0
    url = f"{SUPABASE_URL}/rest/v1/{table}"
    try:
        res = requests.post(url, headers=HEADERS, json=records)
        if res.status_code in [200, 201]:
            return len(records)
        else:
            print(f"  Note on '{table}' ({res.status_code}): {res.text[:120]}")
            return 0
    except Exception as e:
        print(f"  Error inserting into '{table}': {e}")
        return 0

def import_staffing_demands(wb):
    if 'Client_Staffing Log' not in wb.sheetnames:
        return
    print("\n--- 2. Importing Demands (Client_Staffing Log) ---")
    sheet = wb['Client_Staffing Log']
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        return
    
    records = []
    inserted_count = 0
    start_ts = int(time.time())
    for idx, row in enumerate(rows[1:]):
        if not row or not any(row):
            continue
        req_id = f"REQ-{start_ts}-{idx+1:04d}"
        
        client_name = str(row[3]).strip() if len(row) > 3 and row[3] else "Client"
        skill = str(row[6]).strip() if len(row) > 6 and row[6] else "Software Requirement"
        vms_id = str(row[7]).strip() if len(row) > 7 and row[7] else None
        status = str(row[14]).strip() if len(row) > 14 and row[14] else "Open"
        priority = str(row[15]).strip() if len(row) > 15 and row[15] else "Medium"
        
        record = {
            "request_id": req_id,
            "external_requisition_id": vms_id,
            "skill_description": f"{client_name} - {skill}",
            "status": "Open" if "Open" in status else "Closed",
            "priority": "High" if "High" in priority else ("Low" if "Low" in priority else "Medium")
        }
        records.append(record)
        if len(records) >= 200:
            inserted_count += post_batch("demands", records)
            records = []
            
    if records:
        inserted_count += post_batch("demands", records)
    print(f"✅ Demands Migration Completed! ({inserted_count} records)")

def main():
    print("🚀 Running Demand Migration Pipeline...")
    if not os.path.exists(EXCEL_FILE):
        print(f"❌ Excel tracker file not found: {EXCEL_FILE}")
        sys.exit(1)

    print("Loading workbook into memory...")
    wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
    import_staffing_demands(wb)
    print("\n🎉 Demand Migration Completed Successfully!")

if __name__ == '__main__':
    main()
