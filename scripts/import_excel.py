import openpyxl
import os
import sys

# Ensure UTF-8 output encoding for Windows terminal
sys.stdout.reconfigure(encoding='utf-8')

EXCEL_FILE = 'Costaff_Master Tracker_V.1.xlsx'

def main():
    print(f"Starting Excel Data Importer & Inspector for: {EXCEL_FILE}")
    if not os.path.exists(EXCEL_FILE):
        print(f"File not found: {EXCEL_FILE}")
        sys.exit(1)

    size_mb = os.path.getsize(EXCEL_FILE) / (1024 * 1024)
    print(f"File Size: {size_mb:.2f} MB")

    print("Loading workbook sheet names...")
    wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
    sheet_names = wb.sheetnames
    print(f"Discovered {len(sheet_names)} Sheets: {sheet_names}")

    print("\nInspecting Row & Column Counts per Sheet:")
    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        max_r = sheet.max_row
        max_c = sheet.max_column
        print(f"  - [{sheet_name}]: {max_r} rows x {max_c} columns")

    print("\nExcel Data Inspector Completed Successfully!")

if __name__ == '__main__':
    main()
